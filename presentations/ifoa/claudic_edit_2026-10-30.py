"""Claudic edit of David Beddows' IFoA Life Conference deck (20a, 30 Oct 2026).

Typos fixed in place, Dr. added for Jon, a reminder appended to Jon's slide-13 notes,
and PowerPoint (modern) comments on David's slides suggesting what to combine or move.
No slides are removed or reordered.
"""
import re, shutil, sys, uuid, zipfile, os
from pptx import Presentation

SRC, OUT = sys.argv[1], sys.argv[2]
TMP = OUT + ".tmp.pptx"
log = []

TEXT_FIXES = {  # slide number -> [(old, new)]
    # slide 1 "IFoA Life / Conference 2026" is a deliberate line break, not a typo
    2: [("Dr. Chris Martin, Jon Minton", "Dr. Chris Martin, Dr. Jon Minton"),
        ("November 2024", "November 2026"),
        ("long-term term sickness", "long-term sickness")],
    6: [("population and  demographic", "population and demographic")],
    9: [("IIlness", "Illness")],
    23: [("rates : males", "rates: males")],
    25: [("Parker et al.(2020)", "Parker et al. (2020)"),
         ("Study of Aging", "Study of Ageing")],
    27: [("retirements ages", "retirement ages")],
    31: [("Sir Mayfield", "Sir Charlie Mayfield"),
         ("long-term term sickness", "long-term sickness")],
    32: [("one limitations was", "one limitation was")],
}
NOTE_FIXES = {
    18: [("Lievre et al. 2007", "Lièvre et al. 2007")],
    20: [("paramerisation", "parameterisation")],
    33: [("English Longitudinal Survey of Aging", "English Longitudinal Study of Ageing")],
    34: [("English Longitudinal Survey of Aging", "English Longitudinal Study of Ageing")],
}
SLIDE13_NOTE = (
    "REMINDER (Jon): unpack HWLE here, briefly, and say David defines it fully in his section. "
    "From a modelling perspective there are three ways to bound 'working life'. They are conceptually "
    "similar but get thought about slightly differently: "
    "(1) treat retirement as an absorbing state, like death, so time stops accruing once someone retires; "
    "(2) count only the time accumulated in the feeder states (working, sick) on the way to retirement; "
    "(3) only consider paths that start and end within working age (e.g. 50 to 65). "
    "David's model starts everyone working and retires everyone at 65, so it sits closest to (1) and (3)."
)

COMMENTS = {  # slide number -> text
    2: "Typos fixed in this copy: 2024 → 2026 in the date, and 'long-term term sickness' → 'long-term sickness' (also on slide 31). If the programme title still has the double 'term', Giada or Jemma may be able to correct it there too. 'Dr.' added for Jon.",
    3: "The agenda times add up to the full 45 minutes, so the plan fits on paper. Giada's concern is the slide count: about 24 content slides for 35 minutes of talking. The comments on slides 17–28 suggest ways to take 3–4 slides out of the main run without losing any content (moved to the appendix or to notes). If Chris settles on 6 minutes, his row becomes 6 and the spare time could go to Q&A or to David's section.",
    17: "Suggestion: combine 17 and 18. Slide 18 tells the story ('first appears as a term in 2007 … but similar concepts in actuarial journals decades earlier'); 17 adds the actuarial and CMI rows. One option is to keep 17's full timeline under 18's two-part headline and move the reference detail into notes. That saves about a minute and one slide. 17 has about 200 words of labels, the densest slide in the deck.",
    18: "See the comment on 17: a candidate to fold into 17, or to show as a build of the same slide rather than as a separate slide.",
    19: "This overlaps with the Conclusions on 28 (insurance uses, policy uses). Option A: drop 19 from the main run and let 28 carry the message, with 19's bullets moved to notes. Option B: keep 19 and cut 28 back to its three headline statements. Either removes about a minute of repetition.",
    20: "David plans to cover only the key points here. Suggestion: keep the diagram and the advantages and disadvantages, and move the table of rate bases (nn, w, d, r) to the appendix next to slide 32, ready for Q&A.",
    21: "21 and 22 show the same result (13.5 and 0.9 years) as a table and then a chart. Suggestion: keep 22 (the chart) in the main run and move 21 to the appendix. That saves about a minute and one slide. A small style point for a UK audience: thousands show as 100’000 here, and 100,000 may read more naturally.",
    22: "Keep (see the comment on 21).",
    23: "Suggestion: combine 23 and 24. Keep 23's two charts, put 24's three observations underneath them, and move 24's table and worked percentages to notes or the appendix. That saves one slide.",
    24: "The densest results slide (about 180 words). If it stays, suggestion: move the two worked examples (the percentages for a 20-year-old and a 50-year-old) to speaker notes, keep the table, and highlight the age-50 row, since that is the number compared with Parker on 25.",
    25: "Keep: this is the key comparison. Suggestion: put the basis of each number on the slide (13.6 = CMI IP11, occupation class 1, average of females and males, retirement at 65; 12.48 = Parker et al. 2020, healthy and working at 50, from survey data), and move the four bullets to notes. The audience will have seen 13.5 (females) on 21–22, so a label heads off a 'which number?' question.",
    28: "See the comment on 19 about the overlap. If 19 stays, suggestion: keep only the three headline statements here and move the sub-bullets to notes.",
    31: "Typos fixed: 'Sir Mayfield' → 'Sir Charlie Mayfield'; 'long-term term' → 'long-term'.",
}
COMMENT_PREFIX = "Suggestion from Jon (Claude-assisted). "


def iter_text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_text_frames(sh.shapes)
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def replace_in_frame(tf, old, new):
    hits = 0
    for p in tf.paragraphs:
        if old not in p.text:
            continue
        done = False
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done = True
        if not done:  # spans runs: put the paragraph text into the first run
            full = p.text.replace(old, new)
            runs = p.runs
            runs[0].text = full
            for r in runs[1:]:
                r.text = ""
        hits += 1
    return hits


prs = Presentation(SRC)
slides = list(prs.slides)
for n, fixes in TEXT_FIXES.items():
    for old, new in fixes:
        h = sum(replace_in_frame(tf, old, new) for tf in iter_text_frames(slides[n - 1].shapes))
        log.append(f"slide {n}: '{old}' -> '{new}' ({h} hit{'s' if h != 1 else ''})")
for n, fixes in NOTE_FIXES.items():
    tf = slides[n - 1].notes_slide.notes_text_frame
    for old, new in fixes:
        h = replace_in_frame(tf, old, new)
        log.append(f"notes {n}: '{old}' -> '{new}' ({h} hit{'s' if h != 1 else ''})")
tf = slides[12].notes_slide.notes_text_frame
tf.add_paragraph().text = ""
tf.add_paragraph().text = SLIDE13_NOTE
log.append("notes 13: HWLE reminder appended")
prs.save(TMP)

# ---- modern comments, written into the package directly ----
AUTHOR_ID = "{3B42823D-FC15-489F-A24B-AD7901105553}"  # Jon's id from the 13 Oct deck
NS_P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
zin = zipfile.ZipFile(TMP)
files = {i.filename: zin.read(i.filename) for i in zin.infolist()}
zin.close()

authors = files["ppt/authors.xml"].decode()
if AUTHOR_ID not in authors:
    authors = authors.replace("</p188:authorLst>",
        f'<p188:author id="{AUTHOR_ID}" name="Jon Minton" initials="JM" userId="Jon Minton" providerId="None"/></p188:authorLst>')
files["ppt/authors.xml"] = authors.encode()

pres = files["ppt/presentation.xml"].decode()
pres_rels = files["ppt/_rels/presentation.xml.rels"].decode()
rid_to_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="(slides/slide\d+\.xml)"', pres_rels))
rid_to_target.update({k: v for v, k in re.findall(r'Target="(slides/slide\d+\.xml)"[^>]*Id="(rId\d+)"', pres_rels)})
order = re.findall(r'<p:sldId id="(\d+)" r:id="(rId\d+)"', pres)
ct = files["[Content_Types].xml"].decode()

for n, text in COMMENTS.items():
    sld_id, rid = order[n - 1]
    slide_path = "ppt/" + rid_to_target[rid]
    sx = files[slide_path].decode()
    cid = re.search(r'p14:creationId[^>]*val="(\d+)"', sx).group(1)
    name = f"modernComment_{int(sld_id):X}_{int(cid):08X}.xml"
    body = (text if text.startswith("Keep") or text.startswith("See") else text)
    body = COMMENT_PREFIX + body
    body = body.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    files["ppt/comments/" + name] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<p188:cmLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p188="{NS_P188}">'
        f'<p188:cm id="{{{str(uuid.uuid4()).upper()}}}" authorId="{AUTHOR_ID}" created="2026-10-30T15:00:00">'
        '<pc:sldMkLst xmlns:pc="http://schemas.microsoft.com/office/powerpoint/2013/main/command">'
        f'<pc:docMk/><pc:sldMk cId="{cid}" sldId="{sld_id}"/></pc:sldMkLst>'
        f'<p188:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-GB"/><a:t>{body}</a:t></a:r></a:p></p188:txBody>'
        '</p188:cm></p188:cmLst>').encode()
    ct = ct.replace("</Types>",
        f'<Override PartName="/ppt/comments/{name}" ContentType="application/vnd.ms-powerpoint.comments+xml"/></Types>')
    rels_path = slide_path.replace("slides/", "slides/_rels/") + ".rels"
    rels = files[rels_path].decode()
    ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
    new_rid = f"rId{max(ids) + 1}"
    rels = rels.replace("</Relationships>",
        f'<Relationship Id="{new_rid}" Type="http://schemas.microsoft.com/office/2018/10/relationships/comments" Target="../comments/{name}"/></Relationships>')
    files[rels_path] = rels.encode()
    ext = f'<p:ext uri="{{6950BFC3-D8DA-4A85-94F7-54DA5524770B}}"><p188:commentRel xmlns:p188="{NS_P188}" r:id="{new_rid}"/></p:ext>'
    head, tail = sx.split("</p:cSld>", 1)
    if "<p:extLst>" in tail:
        tail = tail.replace("</p:extLst>", ext + "</p:extLst>", 1)
    else:
        tail = tail.replace("</p:sld>", f"<p:extLst>{ext}</p:extLst></p:sld>")
    files[slide_path] = (head + "</p:cSld>" + tail).encode()
    log.append(f"comment on slide {n} ({len(text.split())} words)")

files["[Content_Types].xml"] = ct.encode()
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
    for k, v in files.items():
        z.writestr(k, v)
os.remove(TMP)
print("\n".join(log))
